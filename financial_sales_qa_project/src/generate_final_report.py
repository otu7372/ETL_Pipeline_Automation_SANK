"""
Stage 6b — Generate Final QA Report
----------------------------------------
Consolidates the latest result CSV from each validation stage
(Stage 2 through Stage 5, each produced by its own DefectLogger) into
one Excel workbook with three sheets:

  1. QA Summary   - pass/fail counts per stage, overall pass rate,
                     high-severity defect count, and a "Key Findings"
                     list auto-built from the High severity failures.
  2. Defect Log   - every FAILED check, sorted by severity, with
                     expected/actual/evidence for each - this is the
                     project's core deliverable #7.
  3. Full Results - every check (PASS and FAIL) across all stages,
                     for complete traceability from dashboard number
                     back to source system (deliverable #6).

This script does NOT re-run any validations itself - it only reads
whatever CSVs already exist in reports/. Run src/run_all_stages.py
first (or run each Stage 2-5 script individually) to produce fresh
results before generating the report.

Run from the project root:
    python src/generate_final_report.py
"""

import os
import sys
import glob
import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(PROJECT_ROOT, "reports")

STAGE_PREFIXES = [
    ("Stage2_Source_to_Staging", "Stage 2: Source -> Staging"),
    ("Stage3_Staging_to_Datamart", "Stage 3: Staging -> Data Mart"),
    ("Stage4_Datamart_Aggregates", "Stage 4: Data Mart Aggregates"),
    ("Stage5_Reporting_Layer", "Stage 5: Reporting Layer"),
]

SEVERITY_ORDER = {"High": 0, "Medium": 1, "Low": 2, "-": 3}


def find_latest(prefix: str):
    matches = sorted(glob.glob(os.path.join(REPORTS_DIR, f"{prefix}_*.csv")))
    return matches[-1] if matches else None


def load_all_stage_results() -> pd.DataFrame:
    frames = []
    for prefix, label in STAGE_PREFIXES:
        path = find_latest(prefix)
        if path is None:
            print(f"WARNING: no results found for '{label}'. "
                  f"Run its validation script (or src/run_all_stages.py) first.")
            continue
        df = pd.read_csv(path)
        df["stage_label"] = label
        df["source_file"] = os.path.basename(path)
        frames.append(df)
        print(f"Loaded {len(df)} rows from {os.path.basename(path)} ({label})")

    if not frames:
        raise RuntimeError("No stage result CSVs found in reports/. Run the validation stages first.")

    master = pd.concat(frames, ignore_index=True)
    master["severity_rank"] = master["severity"].map(SEVERITY_ORDER).fillna(4)
    return master


def build_summary_sheet(master: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for prefix, label in STAGE_PREFIXES:
        stage_df = master[master["stage_label"] == label]
        if stage_df.empty:
            continue
        total = len(stage_df)
        passed = (stage_df["status"] == "PASS").sum()
        failed = total - passed
        high = ((stage_df["status"] == "FAIL") & (stage_df["severity"] == "High")).sum()
        rows.append({
            "Stage": label, "Total Checks": total, "Passed": passed, "Failed": failed,
            "High Severity Failures": high,
            "Pass Rate %": round(100 * passed / total, 1) if total else 0
        })

    total_row = {
        "Stage": "OVERALL",
        "Total Checks": len(master),
        "Passed": (master["status"] == "PASS").sum(),
        "Failed": (master["status"] == "FAIL").sum(),
        "High Severity Failures": ((master["status"] == "FAIL") & (master["severity"] == "High")).sum(),
    }
    total_row["Pass Rate %"] = round(100 * total_row["Passed"] / total_row["Total Checks"], 1) \
        if total_row["Total Checks"] else 0
    rows.append(total_row)

    return pd.DataFrame(rows)


def build_key_findings(master: pd.DataFrame) -> pd.DataFrame:
    high_fails = master[(master["status"] == "FAIL") & (master["severity"] == "High")].copy()
    high_fails = high_fails.sort_values(["stage_label", "check_name"])
    cols = ["stage_label", "check_name", "expected", "actual", "details"]
    return high_fails[cols].rename(columns={
        "stage_label": "Stage", "check_name": "Check", "expected": "Expected",
        "actual": "Actual", "details": "Details / Evidence"
    })


def build_defect_log(master: pd.DataFrame) -> pd.DataFrame:
    defects = master[master["status"] == "FAIL"].copy()
    defects = defects.sort_values(["severity_rank", "stage_label", "check_name"])
    cols = ["timestamp", "stage_label", "check_name", "severity", "source_object",
            "target_object", "expected", "actual", "details"]
    out = defects[cols].rename(columns={
        "stage_label": "stage", "check_name": "check_name"
    })
    out.insert(0, "defect_id", [f"DEF-{i+1:03d}" for i in range(len(out))])
    return out


def build_full_results(master: pd.DataFrame) -> pd.DataFrame:
    cols = ["timestamp", "stage_label", "check_name", "status", "severity",
            "source_object", "target_object", "expected", "actual", "details"]
    return master.sort_values(["stage_label", "timestamp"])[cols]


def style_worksheet(ws, header_fill_color="4472C4", freeze=True, autofilter=True):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_fill_color, end_color=header_fill_color, fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=False)

    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        col_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)

    if freeze:
        ws.freeze_panes = "A2"
    if autofilter:
        ws.auto_filter.ref = ws.dimensions


def apply_severity_colors(ws, severity_col_letter):
    fills = {
        "High": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
        "Medium": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
        "Low": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    }
    for row in range(2, ws.max_row + 1):
        cell = ws[f"{severity_col_letter}{row}"]
        fill = fills.get(cell.value)
        if fill:
            for c in ws[row]:
                c.fill = fill


def main():
    print("=" * 70)
    print("STAGE 6: GENERATE FINAL QA REPORT")
    print("=" * 70)

    master = load_all_stage_results()
    summary_df = build_summary_sheet(master)
    key_findings_df = build_key_findings(master)
    defect_log_df = build_defect_log(master)
    full_results_df = build_full_results(master)

    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = os.path.join(REPORTS_DIR, f"QA_Final_Report_{ts}.xlsx")

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="QA Summary", index=False, startrow=0)
        # Key findings appended below the summary table on the same sheet
        start_row = len(summary_df) + 3
        pd.DataFrame([["KEY FINDINGS (High severity)"]]).to_excel(
            writer, sheet_name="QA Summary", index=False, header=False, startrow=start_row
        )
        key_findings_df.to_excel(writer, sheet_name="QA Summary", index=False, startrow=start_row + 1)

        defect_log_df.to_excel(writer, sheet_name="Defect Log", index=False)
        full_results_df.to_excel(writer, sheet_name="Full Results", index=False)

    # Post-process styling
    from openpyxl import load_workbook
    wb = load_workbook(out_path)

    style_worksheet(wb["QA Summary"], freeze=False, autofilter=False)
    style_worksheet(wb["Defect Log"])
    style_worksheet(wb["Full Results"])

    # Color the Defect Log by severity (column E = 'severity' after defect_id/timestamp/stage/check_name)
    defect_cols = list(defect_log_df.columns)
    severity_col_idx = defect_cols.index("severity") + 1
    severity_col_letter = get_column_letter(severity_col_idx)
    apply_severity_colors(wb["Defect Log"], severity_col_letter)

    wb.save(out_path)

    print(f"\nFinal QA report written to: {out_path}")
    print(f"  - QA Summary sheet: {len(summary_df)} stage rows + {len(key_findings_df)} key findings")
    print(f"  - Defect Log sheet: {len(defect_log_df)} defects")
    print(f"  - Full Results sheet: {len(full_results_df)} total checks")


if __name__ == "__main__":
    main()
